"""Turn documents into markdown text with location markers (pages, slides, transcript times)."""

from __future__ import annotations

import html
import re
from html.parser import HTMLParser
from pathlib import Path


class ParseError(Exception):
    pass


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    pages = []
    for number, page in enumerate(PdfReader(path).pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"<!-- page {number} -->\n{text}")
    if not pages:
        raise ParseError("no extractable text (scanned PDF? OCR is not supported yet)")
    return "\n\n".join(pages)


def _table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    cells = [[c.replace("\n", " ").replace("|", "\\|").strip() for c in row] for row in rows]
    lines = ["| " + " | ".join(cells[0]) + " |", "|" + "---|" * len(cells[0])]
    lines += ["| " + " | ".join(row) + " |" for row in cells[1:]]
    return "\n".join(lines)


def _docx(path: Path) -> str:
    import docx
    from docx.table import Table

    out = []
    for block in docx.Document(str(path)).iter_inner_content():
        if isinstance(block, Table):
            out.append(_table([[cell.text for cell in row.cells] for row in block.rows]))
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        if style.startswith("Heading") and style[-1:].isdigit():
            out.append("#" * int(style[-1]) + " " + text)
        elif style == "Title":
            out.append("# " + text)
        else:
            out.append(text)
    return "\n\n".join(out)


def _pptx(path: Path) -> str:
    from pptx import Presentation

    slides = []
    for number, slide in enumerate(Presentation(str(path)).slides, 1):
        parts = [f"<!-- slide {number} -->"]
        title = slide.shapes.title
        if title is not None and title.has_text_frame and title.text_frame.text.strip():
            parts.append("## " + title.text_frame.text.strip())
        for shape in slide.shapes:
            if shape == title:
                continue
            if getattr(shape, "has_table", False):
                parts.append(_table([[c.text for c in row.cells] for row in shape.table.rows]))  # type: ignore[attr-defined]
            elif shape.has_text_frame and shape.text_frame.text.strip():  # pyright: ignore[reportAttributeAccessIssue]
                parts.append(shape.text_frame.text.strip())  # pyright: ignore[reportAttributeAccessIssue]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        slides.append("\n\n".join(parts))
    return "\n\n".join(slides)


class _HTML(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "table", "ul", "ol", "header", "footer"}

    def __init__(self) -> None:
        super().__init__()
        self.out: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.out.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag == "li":
            self.out.append("\n- ")
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "noscript"):
            self.skip -= 1
        elif tag in self.BLOCK or tag.startswith("h"):
            self.out.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skip:
            self.out.append(" ".join(data.split()) if data.strip() else "")


def _html(path: Path) -> str:
    parser = _HTML()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    lines = [line.strip() for line in "".join(parser.out).splitlines()]
    return "\n".join(line for i, line in enumerate(lines) if line or (i and lines[i - 1])).strip()


CUE_TIME_RE = re.compile(r"^(?:(\d+):)?(\d{1,2}):(\d{2})[.,]\d{1,3}\s+-->")
VOICE_RE = re.compile(r"<v(?:\.[^\s>]*)?\s+([^>]+)>")
SPEAKER_RE = re.compile(r"^([^\W\d][\w.'-]*(?: [\w.'-]+){0,3}):\s+(.+)$")  # "Alice Smith: text"
PARAGRAPH_SECONDS = 60


def _transcript(path: Path) -> str:
    """WebVTT or SRT → one line per speaker turn, `[hh:mm:ss] Speaker: text`. Consecutive cues of the
    same speaker are joined, with a fresh timestamp at least every minute so claims stay locatable."""
    turns: list[list] = []  # [seconds, speaker, text]
    for block in re.split(r"\n\s*\n", path.read_text(encoding="utf-8-sig", errors="replace").replace("\r", "")):
        lines = block.strip().splitlines()
        timing = next((i for i, line in enumerate(lines) if "-->" in line), None)
        if timing is None:  # header, NOTE, STYLE, REGION
            continue
        m = CUE_TIME_RE.match(lines[timing].strip())
        if not m:
            continue
        seconds = int(m.group(1) or 0) * 3600 + int(m.group(2)) * 60 + int(m.group(3))
        raw = " ".join(lines[timing + 1:])
        voice = VOICE_RE.search(raw)
        text = " ".join(html.unescape(re.sub(r"<[^>]+>", "", raw)).split())
        speaker = voice.group(1).strip() if voice else None
        if not speaker and (named := SPEAKER_RE.match(text)):
            speaker, text = named.groups()
        if not text:
            continue
        last = turns[-1] if turns else None
        if last and last[1] == speaker and seconds - last[0] < PARAGRAPH_SECONDS:
            last[2] += " " + text
        else:
            turns.append([seconds, speaker, text])
    if not turns:
        raise ParseError("no transcript cues found")
    stamp = lambda t: f"{t // 3600:02d}:{t % 3600 // 60:02d}:{t % 60:02d}"  # noqa: E731
    return "<!-- transcript: [hh:mm:ss] is the time from the start of the recording -->\n\n" + "\n\n".join(
        f"[{stamp(t)}] {speaker + ': ' if speaker else ''}{text}" for t, speaker, text in turns)


PARSERS = {
    ".md": lambda p: p.read_text(encoding="utf-8", errors="replace"),
    ".markdown": lambda p: p.read_text(encoding="utf-8", errors="replace"),
    ".txt": lambda p: p.read_text(encoding="utf-8", errors="replace"),
    ".pdf": _pdf,
    ".docx": _docx,
    ".pptx": _pptx,
    ".html": _html,
    ".htm": _html,
    ".vtt": _transcript,
    ".srt": _transcript,
}


def parse(path: Path) -> str:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ParseError(f"unsupported file type '{path.suffix}' (supported: {', '.join(sorted(PARSERS))})")
    try:
        text = parser(path).strip()
    except ParseError:
        raise
    except Exception as exc:  # third-party parsers raise anything on broken files
        raise ParseError(f"could not read {path.name}: {exc}") from exc
    if not text:
        raise ParseError(f"{path.name} contains no text")
    return text
