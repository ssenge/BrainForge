"""Slices 3, 5, 6, 7: ingestion pipeline, idempotency, guards, formats, inbox."""

import pytest
from conftest import commits, make_pdf, write

from brainforge import agents, gitops, ingest, okf, parsers


def cfg(brain):
    return agents.resolve(brain)


def snapshot(brain):
    return gitops.head(brain.root), gitops.status(brain.root)


def test_cli_ingest_creates_cited_concepts_archives_and_commits(brain, tmp_path, run_cli):
    doc = write(tmp_path / "notes.md", "Topic: Quantum annealing\nTopic: QAOA\n")
    code, out, err = run_cli("--brain", brain.root, "ingest", doc)
    assert code == 0, out + err
    topic = okf.read(brain.knowledge / "topics" / "qaoa.md")
    source = next((brain.knowledge / "sources").glob("notes-*.md"))
    assert topic.meta["sources"][0]["resource"] == f"../sources/{source.name}"
    assert topic.meta["generated"]["by"] == "brainforge/command"
    src = okf.read(source)
    assert src.meta["brainforge"]["sha256"] == ingest.sha256(doc)
    assert src.meta["brainforge"]["entry"] == "cli"
    assert src.meta["title"] == "Test document"
    assert "[qaoa](../topics/qaoa.md) (created)" in src.body
    archived = brain.root / src.meta["resource"].removeprefix("../../")
    assert archived.read_bytes() == doc.read_bytes() and doc.exists()  # copied, original untouched
    assert commits(brain.root)[0] == "ingest: notes.md"
    assert "BrainForge-Source: knowledge/sources/" in gitops.git(brain.root, "log", "-1", "--format=%B")
    assert gitops.status(brain.root) == []
    assert "(qaoa.md) - Topic QAOA." in (brain.knowledge / "topics" / "index.md").read_text()
    assert "**Ingestion**: [notes.md]" in (brain.knowledge / "log.md").read_text()
    assert [i for i in okf.validate_bundle(brain.knowledge) if i.level == "error"] == []


def test_reingesting_same_content_is_a_noop(brain, tmp_path):
    doc = write(tmp_path / "a.md", "Topic: Alpha\n")
    assert ingest.ingest(brain, doc, cfg(brain)).status == "ingested"
    head = gitops.head(brain.root)
    again = ingest.ingest(brain, doc, cfg(brain))
    renamed = ingest.ingest(brain, write(tmp_path / "renamed.md", "Topic: Alpha\n"), cfg(brain))
    assert (again.status, renamed.status) == ("skipped", "skipped")
    assert gitops.head(brain.root) == head


def test_force_reuses_the_same_source_concept(brain, tmp_path):
    doc = write(tmp_path / "a.md", "Topic: Alpha\n")
    first = ingest.ingest(brain, doc, cfg(brain))
    forced = ingest.ingest(brain, doc, cfg(brain), force=True)
    assert forced.status == "ingested" and forced.source == first.source
    assert len(okf.concept_files(brain.knowledge / "sources")) == 1


@pytest.mark.parametrize("mode,reason", [
    ("outside", "outside knowledge/"),
    ("delete", "deleted"),
    ("invalid", "missing YAML frontmatter"),
    ("nocite", "does not cite"),
    ("drop-sources", "removed existing sources"),
    ("fail", "exited with 3"),
])
def test_rejected_agent_output_leaves_the_brain_untouched(brain, tmp_path, monkeypatch, mode, reason):
    ingest.ingest(brain, write(tmp_path / "seed.md", "Topic: Seed\n"), cfg(brain))
    before = snapshot(brain)
    monkeypatch.setenv("FAKE_AGENT", mode)
    outcome = ingest.ingest(brain, write(tmp_path / "doc.md", "Topic: Beta\n"), cfg(brain))
    assert outcome.status == "failed" and reason in outcome.message
    assert snapshot(brain) == before


def test_dry_run_validates_but_changes_nothing(brain, tmp_path):
    before = snapshot(brain)
    outcome = ingest.ingest(brain, write(tmp_path / "d.md", "Topic: Gamma\n"), cfg(brain), dry_run=True)
    assert outcome.status == "dry-run" and outcome.created == ["knowledge/topics/gamma.md"]
    assert snapshot(brain) == before


def test_dirty_tree_is_refused(brain, tmp_path, run_cli):
    (brain.root / "instructions.md").write_text("changed\n")
    code, _, err = run_cli("--brain", brain.root, "ingest", write(tmp_path / "d.md", "Topic: X\n"))
    assert code == 2 and "uncommitted changes" in err


def test_updates_preserve_history_and_provenance(brain, tmp_path, run_cli):
    """Journey 4: a rescheduled milestone updates the same concept and keeps both sources."""
    ingest.ingest(brain, write(tmp_path / "plan.md", "Milestone: Integration testing = 2026-11-10\n"), cfg(brain))
    out = ingest.ingest(brain, write(tmp_path / "meeting.md", "Milestone: Integration testing = 2026-11-17\n"), cfg(brain))
    assert out.updated == ["knowledge/milestones/integration-testing.md"] and out.created == []
    milestone = okf.read(brain.knowledge / "milestones" / "integration-testing.md")
    assert len(milestone.sources()) == 2
    assert "2026-11-10" in milestone.body and "2026-11-17" in milestone.body
    assert len(list((brain.knowledge / "milestones").glob("*.md"))) == 2  # concept + index


def test_instruction_layers_reach_the_agent_in_order(brain, tmp_path, prompts):
    write(brain.root / "instructions" / "ingestion.md", "BRAIN-INGESTION-RULE\n")
    gitops.commit(brain.root, "add instructions", ["instructions"])
    doc = write(tmp_path / "x.pdf.md", "Topic: X\n")
    write(tmp_path / "x.pdf.md.sidecar.md", "---\ndocument: x.pdf.md\n---\nSIDECAR-FOCUS\n")
    ingest.ingest(brain, doc, cfg(brain), instruction="CLI-FOCUS")
    p = prompts()
    order = [p.index(s) for s in ("BrainForge framework policy", "Ingestion: extract", "BRAIN-INGESTION-RULE",
                                  "SIDECAR-FOCUS", "CLI-FOCUS", "<document name=")]
    assert order == sorted(order)
    source = okf.read(okf.concept_files(brain.knowledge / "sources")[0])
    kinds = [s.split(":")[0].split("@")[0] for s in source.meta["brainforge"]["instructions"]]
    assert kinds == ["framework", "framework", "instructions.md", "instructions/ingestion.md", "sidecar", "cli"]


def test_inbox_processing_with_sidecar_trailer_and_duplicates(brain, tmp_path, prompts):
    write(brain.inbox / "minutes.md", "Topic: Delivery risk\n")
    write(brain.inbox / "minutes.md.sidecar.md", "Focus on risks.\n")
    write(brain.inbox / "deck.md", "Topic: Architecture\n")
    gitops.git(brain.root, "add", "inbox")
    gitops.git(brain.root, "commit", "-q", "-m", "ingest: workshop\n\nIngest-Note: Focus on milestone changes")
    write(brain.inbox / "orphan.pdf.sidecar.md", "lonely\n")

    outcomes, orphans = ingest.process_inbox(brain, cfg(brain))
    assert [o.status for o in outcomes] == ["ingested", "ingested"]
    assert [p.name for p in orphans] == ["orphan.pdf.sidecar.md"]
    assert "Focus on milestone changes" in prompts() and "Focus on risks." in prompts()
    assert sorted(p.name for p in brain.inbox.iterdir()) == [".gitkeep", "orphan.pdf.sidecar.md"]
    archived = sorted(p.name for p in brain.archive.rglob("*.md"))
    assert archived == ["deck.md", "minutes.md", "minutes.md.sidecar.md"]
    assert commits(brain.root)[:2] == ["ingest: minutes.md", "ingest: deck.md"]
    assert okf.read(next((brain.knowledge / "sources").glob("deck-*.md"))).meta["brainforge"]["entry"] == "inbox"

    # The same bytes pushed again: removed from the inbox, nothing re-ingested.
    write(brain.inbox / "deck-copy.md", "Topic: Architecture\n")
    gitops.commit(brain.root, "ingest: again", ["inbox"])
    outcomes, _ = ingest.process_inbox(brain, cfg(brain))
    assert [o.status for o in outcomes] == ["skipped"]
    assert not (brain.inbox / "deck-copy.md").exists()
    assert commits(brain.root)[0].startswith("inbox: remove deck-copy.md")


def test_failed_inbox_document_stays_for_retry(brain, monkeypatch):
    write(brain.inbox / "doc.md", "Topic: Retry me\n")
    gitops.commit(brain.root, "add", ["inbox"])
    monkeypatch.setenv("FAKE_AGENT", "fail")
    outcomes, _ = ingest.process_inbox(brain, cfg(brain))
    assert outcomes[0].status == "failed" and (brain.inbox / "doc.md").exists()
    monkeypatch.setenv("FAKE_AGENT", "good")
    outcomes, _ = ingest.process_inbox(brain, cfg(brain))
    assert outcomes[0].status == "ingested" and not (brain.inbox / "doc.md").exists()


def test_parsers_keep_locations(tmp_path):
    import docx
    import pptx

    pdf = parsers.parse(make_pdf(tmp_path / "p.pdf", "Topic: Quantum"))
    assert pdf.startswith("<!-- page 1 -->") and "Topic: Quantum" in pdf

    d = docx.Document()
    d.add_heading("Plan", level=1)
    d.add_paragraph("Integration testing on 2026-11-17.")
    table = d.add_table(rows=2, cols=2)
    for r, row in enumerate([["Owner", "Task"], ["Alice", "Test"]]):
        for col, value in enumerate(row):
            table.cell(r, col).text = value
    d.save(tmp_path / "d.docx")
    text = parsers.parse(tmp_path / "d.docx")
    assert "# Plan" in text and "| Alice | Test |" in text

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Go-live in March"
    deck.save(tmp_path / "s.pptx")
    text = parsers.parse(tmp_path / "s.pptx")
    assert text.startswith("<!-- slide 1 -->") and "## Roadmap" in text and "Go-live in March" in text

    html = write(tmp_path / "h.html", "<html><script>x()</script><h2>Risks</h2><ul><li>Delay</li></ul></html>")
    assert parsers.parse(html) == "## Risks\n\n- Delay"

    vtt = write(tmp_path / "t.vtt", "WEBVTT\n\nNOTE x\n\n1\n00:00:01.000 --> 00:00:03.000 align:start\n"
                "<v Alice Smith>Go-live</v>\n\n00:00:03.500 --> 00:00:05.000\n<v Alice Smith>moves &amp; slips.</v>\n\n"
                "01:02:03.000 --> 01:02:04.000\n<v.loud Bob>OK</v>\n\n01:03:04.000 --> 01:03:05.000\nBob: Agreed.\n")
    assert parsers.parse(vtt).splitlines()[2:] == [
        "[00:00:01] Alice Smith: Go-live moves & slips.", "", "[01:02:03] Bob: OK", "", "[01:03:04] Bob: Agreed."]
    srt = write(tmp_path / "t.srt", "1\r\n00:00:01,000 --> 00:00:02,000\r\nAlice: Risk: vendor\r\nlate.\r\n\r\n"
                "2\r\n00:00:02,500 --> 00:00:03,000\r\nNo speaker.\r\n")
    assert parsers.parse(srt).endswith("[00:00:01] Alice: Risk: vendor late.\n\n[00:00:02] No speaker.")
    with pytest.raises(parsers.ParseError, match="no transcript cues"):
        parsers.parse(write(tmp_path / "e.vtt", "WEBVTT\n\nNOTE nothing\n"))


def test_unsupported_and_empty_documents_fail_cleanly(brain, tmp_path):
    with pytest.raises(parsers.ParseError, match="unsupported"):
        parsers.parse(write(tmp_path / "x.xyz", "data"))
    blank = make_pdf(tmp_path / "blank.pdf", "")
    out = ingest.ingest(brain, blank, cfg(brain))
    assert out.status == "failed" and "no extractable text" in out.message
